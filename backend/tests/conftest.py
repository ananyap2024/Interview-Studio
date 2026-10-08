from io import BytesIO

import pymupdf as fitz
import pytest
from docx import Document


@pytest.fixture
def pdf_bytes() -> bytes:
    document = fitz.open()
    page = document.new_page()
    page.insert_text((72, 72), "Candidate Name\nPython Developer")
    content = document.tobytes()
    document.close()
    return content


@pytest.fixture
def empty_pdf_bytes() -> bytes:
    document = fitz.open()
    document.new_page()
    content = document.tobytes()
    document.close()
    return content


@pytest.fixture
def docx_bytes() -> bytes:
    document = Document()
    document.add_heading("Candidate Name", level=1)
    document.add_paragraph("Python Developer")
    output = BytesIO()
    document.save(output)
    return output.getvalue()